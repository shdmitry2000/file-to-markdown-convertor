"""Two LLM modes for every converter call: the space's connector, or this pod's lane.

space  — the space's connector only; the pod's env is never read, and a connector
         missing what its type needs fails the job before any work, naming the field.
params — this pod's lane (VLM_MODEL / LLM_*), with request fields on top.

The pod here is OCP's markdown pod: proxy lane configured, no PROXY_AUTH_TOKEN.
Real files go through the real worker (cold and warm); the LLM itself is a stub.
"""
from __future__ import annotations

import io
import logging
import queue
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from PIL import Image

from app.converters import vlm_lane
from app.workers import worker

SPACE = {"mode": "space", "model": "vertex_proxy/gemini-3.5-flash",
         "base_url": "https://space-proxy", "connector_id": "space-c",
         "auth_token": "SPACE-TOKEN", "project": "p"}


@pytest.fixture(autouse=True)
def ocp_markdown_pod(monkeypatch):
    for k in ("PROXY_AUTH_TOKEN", "VLM_MODEL", "LLM_FACTORY_PROVIDER"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("LLM_CONNECTOR_TYPE", "vertex_proxy")
    monkeypatch.setenv("LLM_TYPE", "gemini-3.8-flash")
    monkeypatch.setenv("PROXY_BASE_URL", "https://pod-proxy")
    monkeypatch.setenv("PROXY_CONNECTOR_ID", "pod-c")
    vlm_lane._clients.clear()


def _vision(text: str) -> MagicMock:
    c = MagicMock()
    c.vision_sync.return_value = text
    return c


# --- space mode ---------------------------------------------------------------------

def test_space_mode_builds_the_spaces_own_client_and_model():
    from shared.llm_factory.backends.litellm.vertex_proxy import VertexProxyClient

    with vlm_lane.job(SPACE):
        c = vlm_lane.client()
    assert isinstance(c, VertexProxyClient)
    assert (c._base_url, c._connector_id, c._auth_token) == (
        "https://space-proxy", "space-c", "SPACE-TOKEN")
    assert c.provider == "vertex_proxy/gemini-3.5-flash"  # the space's model, not the pod's


def test_space_mode_never_borrows_the_pods_token(monkeypatch):
    monkeypatch.setenv("PROXY_AUTH_TOKEN", "POD-TOKEN")
    with pytest.raises(vlm_lane.LaneError, match="has no auth_token"):
        with vlm_lane.job({**SPACE, "auth_token": ""}):
            pass


@pytest.mark.parametrize("drop, named", [
    (("auth_token",), "auth_token"),
    (("base_url", "connector_id"), "base_url, connector_id"),
])
def test_an_incomplete_space_connector_names_what_is_missing(drop, named):
    conn = {k: v for k, v in SPACE.items() if k not in drop}
    with pytest.raises(vlm_lane.LaneError, match=named):
        vlm_lane.check(conn)


def test_a_space_connector_without_a_model_fails():
    with pytest.raises(vlm_lane.LaneError, match="no model"):
        vlm_lane.check({**SPACE, "model": ""})


def test_a_vertex_ai_space_needs_a_project_but_not_a_location():
    vlm_lane.check({"mode": "space", "model": "vertex_ai/gemini-3.5-flash", "project": "p"})
    with pytest.raises(vlm_lane.LaneError, match="project"):
        vlm_lane.check({"mode": "space", "model": "vertex_ai/gemini-3.5-flash"})


def test_space_mode_escalates_on_the_spaces_lane():
    with patch("shared.space_llm.pinned_client", return_value=_vision("x")) as pin, \
         vlm_lane.job(SPACE):
        vlm_lane.client("gemini-9-pro")
        assert vlm_lane.model("gemini-9-pro") == "vertex_proxy/gemini-9-pro"
    creds, spec = pin.call_args.args
    assert spec == "vertex_proxy/gemini-9-pro" and creds.auth_token == "SPACE-TOKEN"


# --- params mode --------------------------------------------------------------------

def test_params_mode_with_no_fields_is_the_pods_lane():
    with patch("shared.llm_factory.LLMFactory.build", return_value=_vision("x")) as build:
        with vlm_lane.job({"mode": "params"}):
            vlm_lane.client()
        vlm_lane.client()  # no llm at all: the same lane
    assert [c.args for c in build.call_args_list] == [("vertex_proxy/gemini-3.8-flash",)]


def test_params_mode_request_model_overrides_the_pods(monkeypatch):
    monkeypatch.setenv("VLM_MODEL", "vertex_ai/gemini-3.5-flash-lite")
    with patch("shared.llm_factory.LLMFactory.build", return_value=_vision("x")) as build, \
         vlm_lane.job({"mode": "params", "model": "vertex_ai/gemini-3.8-flash"}):
        vlm_lane.client()
    build.assert_called_once_with("vertex_ai/gemini-3.8-flash")


def test_params_mode_request_identity_overrides_the_pods():
    from shared.llm_factory.backends.litellm.vertex_proxy import VertexProxyClient

    conn = {"mode": "params", "model": "vertex_proxy/gemini-3.8-flash",
            "base_url": "https://req", "connector_id": "req-c", "auth_token": "REQ-TOKEN"}
    with vlm_lane.job(conn):
        c = vlm_lane.client()
    assert isinstance(c, VertexProxyClient) and c._auth_token == "REQ-TOKEN"


def test_a_lane_without_vision_is_refused():
    with patch("shared.llm_factory.LLMFactory.build", return_value=object()), \
         pytest.raises(vlm_lane.LaneError, match="vision"):
        vlm_lane.client()


# --- the API ------------------------------------------------------------------------

@pytest.fixture
def api(tmp_path, monkeypatch):
    from fastapi.testclient import TestClient

    from tests.zmq_fakes import FakeContext
    monkeypatch.setattr("zmq.Context", FakeContext)
    from app.api import main

    sent = []
    monkeypatch.setattr(main, "enqueue_task", sent.append)
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    with TestClient(main.app) as c:
        yield c, sent, str(pdf)


def test_the_api_sends_the_mode_and_never_logs_the_token(api, caplog):
    c, sent, pdf = api
    caplog.set_level(logging.DEBUG)
    assert c.post("/convert", json={"file_path": pdf, "llm": SPACE}).status_code == 200
    assert sent[-1]["llm_connector"] == SPACE
    assert "SPACE-TOKEN" not in caplog.text


def test_the_api_defaults_llm_to_params_mode(api):
    c, sent, pdf = api
    assert c.post("/convert", json={"file_path": pdf, "llm": {"model": "vertex_ai/x"}}).status_code == 200
    assert sent[-1]["llm_connector"] == {"model": "vertex_ai/x", "mode": "params"}


def test_the_api_refuses_an_unknown_mode_before_queueing(api):
    c, sent, pdf = api
    r = c.post("/convert", json={"file_path": pdf, "llm": {"mode": "pod"}})
    assert r.status_code == 400 and "mode" in r.json()["detail"] and sent == []


def test_a_pre_1_4_9_connector_passes_through_without_a_mode(api):
    c, sent, pdf = api
    old = {k: v for k, v in SPACE.items() if k != "mode"}
    assert c.post("/convert", json={"file_path": pdf, "llm_connector": old}).status_code == 200
    assert sent[-1]["llm_connector"] == old


# --- real files through the worker --------------------------------------------------

def _png() -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (300, 300), (200, 30, 30)).save(buf, format="PNG")
    buf.seek(0)
    return buf


def _docx_with_picture(path: Path) -> Path:
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.add_paragraph("דוח תחקיר אירוע")
    d.add_picture(_png(), width=Inches(2))
    d.save(str(path))
    return path


def _slide_pdf(path: Path) -> Path:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=_png().getvalue())
    doc.save(str(path))
    return path


def _cold(path: Path, converter: str, conn):
    q = queue.Queue()
    worker._convert_in_subprocess(str(path), converter, False, q, True, conn)
    return q.get_nowait()


def _warm(path: Path, conn):
    in_q, out_q = queue.Queue(), queue.Queue()
    in_q.put((str(path), conn))
    in_q.put(None)
    worker._warm_docling_loop(in_q, out_q, False, True)
    replies = [out_q.get_nowait() for _ in range(out_q.qsize())]
    assert replies[0][0] == "ready", replies[0]
    return replies[1]


INCOMPLETE = {**SPACE, "auth_token": ""}


@pytest.mark.parametrize("converter", ["docling", "dbank", "slides", "vlm"])
def test_an_incomplete_space_fails_the_cold_job_before_any_work(tmp_path, converter):
    with patch.object(worker, "_convert_job") as work:
        status, detail = _cold(tmp_path / "x.pdf", converter, INCOMPLETE)[:2]
    assert status == "error" and "has no auth_token" in detail
    work.assert_not_called()


def test_an_incomplete_space_fails_the_warm_job(tmp_path):
    status, detail = _warm(_docx_with_picture(tmp_path / "a.docx"), INCOMPLETE)[:2]
    assert status == "error" and "has no auth_token" in detail


@pytest.mark.parametrize("path", ["cold", "warm"])
def test_space_mode_describes_docx_pictures_with_the_spaces_client(tmp_path, path):
    space_client = _vision("ריבוע אדום")
    doc = _docx_with_picture(tmp_path / "a.docx")
    with patch("shared.space_llm.pinned_client", return_value=space_client) as pin, \
         patch("shared.llm_factory.LLMFactory.build") as pod:
        reply = _cold(doc, "docling", SPACE) if path == "cold" else _warm(doc, SPACE)
    status, payload = reply[:2]
    assert status == "success", payload
    assert "ריבוע אדום" in payload["markdown"]
    assert pin.call_args.args[1] == "vertex_proxy/gemini-3.5-flash"
    pod.assert_not_called()


@pytest.mark.parametrize("converter", ["slides", "vlm"])
def test_space_mode_runs_page_converters_on_the_spaces_client(tmp_path, converter):
    space_client = _vision("טקסט השקף")
    with patch("shared.space_llm.pinned_client", return_value=space_client), \
         patch("shared.llm_factory.LLMFactory.build") as pod:
        status, payload = _cold(_slide_pdf(tmp_path / "deck.pdf"), converter, SPACE)[:2]
    assert status == "success", payload
    assert "טקסט השקף" in payload["markdown"]
    assert space_client.vision_sync.called
    pod.assert_not_called()


def test_params_mode_on_this_pod_fails_page_converters_with_the_pods_reason(tmp_path):
    """Params mode uses the pod's lane, so OCP's missing token is what the user sees."""
    status, detail = _cold(_slide_pdf(tmp_path / "deck.pdf"), "vlm", {"mode": "params"})[:2]
    assert status == "error" and "PROXY_AUTH_TOKEN" in detail
