"""OCP's markdown pod: proxy lane configured, but no PROXY_AUTH_TOKEN of its own.

Real docx files go through the real docling worker (cold and warm paths), the way
investigation forms did on OCP. They used to fail the whole file with "vertex_proxy:
missing required env var PROXY_AUTH_TOKEN": the table-heavy form read as unusable,
went to the VLM fallback, and the fallback built the pod's own client. No network:
any VLM call in these tests would fail to build its client, which is the point.
"""
from __future__ import annotations

import io
import queue
from pathlib import Path

import pytest
from PIL import Image

from app.workers import worker

TYPE_ONLY = {"model": "vertex_proxy/gemini-3.5-flash"}  # a 1.4.7.8-1.4.7.9 template


@pytest.fixture(autouse=True)
def ocp_markdown_pod(monkeypatch):
    for k in ("PROXY_AUTH_TOKEN", "VLM_MODEL", "LLM_FACTORY_PROVIDER"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("LLM_CONNECTOR_TYPE", "vertex_proxy")
    monkeypatch.setenv("LLM_TYPE", "gemini-3.8-flash")
    monkeypatch.setenv("PROXY_BASE_URL", "https://proxy.invalid")
    monkeypatch.setenv("PROXY_CONNECTOR_ID", "c-1")
    monkeypatch.setenv("VLM_FALLBACK_ENABLED", "true")
    from app.converters import vlm_lane
    vlm_lane._clients.clear()


def _form_docx(path: Path, picture: bool = False) -> Path:
    """An investigation form: a big table, mostly empty cells, little text."""
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.add_paragraph("טופס תחקיר")
    t = d.add_table(rows=60, cols=6)
    t.cell(0, 0).text = "שם"
    t.cell(1, 0).text = "תאריך"
    if picture:
        buf = io.BytesIO()
        Image.new("RGB", (300, 300), (200, 30, 30)).save(buf, format="PNG")
        buf.seek(0)
        d.add_picture(buf, width=Inches(2))
    d.save(str(path))
    return path


def test_the_pod_really_cannot_build_its_own_lane():
    # Precondition: this env reproduces OCP, or the tests below prove nothing.
    from shared.llm_factory import LLMFactory
    with pytest.raises(ValueError, match="PROXY_AUTH_TOKEN"):
        LLMFactory.from_env()


def _cold(path: Path, connector=None):
    q = queue.Queue()
    worker._convert_in_subprocess(str(path), "docling", False, q, True, connector)
    return q.get_nowait()


@pytest.mark.parametrize("connector", [None, TYPE_ONLY])
@pytest.mark.parametrize("picture", [False, True])
def test_a_table_form_converts_on_the_cold_path(tmp_path, connector, picture):
    status, payload = _cold(_form_docx(tmp_path / "301_טופס.docx", picture), connector)[:2]
    assert status == "success", payload
    assert payload["converter_used"] == "docling"
    assert "תאריך" in payload["markdown"]
    assert worker._docling_output_unusable(payload["markdown"])  # it IS the risky shape


@pytest.mark.parametrize("connector", [None, TYPE_ONLY])
def test_a_table_form_converts_on_the_warm_path(tmp_path, connector):
    in_q, out_q = queue.Queue(), queue.Queue()
    in_q.put((str(_form_docx(tmp_path / "376_תחקיר.docx", picture=True)), connector))
    in_q.put(None)
    worker._warm_docling_loop(in_q, out_q, False, True)
    replies = [out_q.get_nowait() for _ in range(out_q.qsize())]
    assert replies[0][0] == "ready", replies[0]
    status, payload = replies[1][:2]
    assert status == "success", payload
    assert payload["converter_used"] == "docling"
    assert "תאריך" in payload["markdown"]
