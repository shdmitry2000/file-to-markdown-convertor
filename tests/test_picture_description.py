"""Pictures inside docx/pptx are described by our VLM (mocked — no network)."""
from __future__ import annotations

import io
from pathlib import Path
from unittest.mock import patch

import pytest
from PIL import Image

from app.converters import picture_description as pd
from app.converters.image_ids import annotate_images


def _png(side: int) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", (side, side), (200, 30, 30)).save(buf, format="PNG")
    buf.seek(0)
    return buf


def _docx(path: Path, side: int = 200, in_table: bool = False) -> Path:
    import docx
    from docx.shared import Inches

    d = docx.Document()
    d.add_paragraph("שלום עולם")
    if in_table:
        cell = d.add_table(rows=1, cols=2).cell(0, 0)
        cell.paragraphs[0].add_run().add_picture(_png(side), width=Inches(1))
    else:
        d.add_picture(_png(side), width=Inches(1))
    d.save(str(path))
    return path


def _pptx(path: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    slide.shapes.add_picture(_png(600), Inches(1), Inches(1), Inches(6), Inches(4))
    p.save(str(path))
    return path


def _convert(path: Path) -> str:
    from docling.document_converter import DocumentConverter

    result = DocumentConverter(format_options=pd.office_format_options()).convert(str(path))
    return annotate_images(result.document.export_to_markdown(), result.document)


@pytest.fixture
def vlm():
    with patch("app.converters.vlm.VLMConverter.__init__", return_value=None), \
         patch("app.converters.vlm.VLMConverter.transcribe_page") as tp:
        yield tp


def test_docx_picture_gets_its_description_after_the_tag(tmp_path, vlm):
    vlm.return_value = "גרף רווחים\n2024: 1,200"
    md = _convert(_docx(tmp_path / "a.docx"))
    assert "<!-- image page=" in md  # image identity still applies
    assert "גרף רווחים 2024: 1,200" in md  # one line, Hebrew kept
    assert vlm.call_count == 1


def test_description_in_a_table_cell_keeps_the_table(tmp_path, vlm):
    vlm.return_value = "a | b\nc"
    md = _convert(_docx(tmp_path / "a.docx", in_table=True))
    row = next(line for line in md.splitlines() if "a / b c" in line)
    assert row.count("|") == 3  # two cells, not broken by the description


def test_pptx_picture_is_described(tmp_path, vlm):
    vlm.return_value = "מפת סניפים"
    assert "מפת סניפים" in _convert(_pptx(tmp_path / "a.pptx"))


def test_a_failed_description_does_not_fail_the_document(tmp_path, vlm):
    vlm.side_effect = RuntimeError("blocked")
    md = _convert(_docx(tmp_path / "a.docx"))
    assert "שלום עולם" in md and "<!-- image" in md


def test_icons_are_not_sent_to_the_vlm(tmp_path, vlm):
    _convert(_docx(tmp_path / "a.docx", side=16))
    vlm.assert_not_called()


@pytest.mark.parametrize("fmt", ["DOCX", "PPTX"])
def test_docling_pipeline_finds_our_engine(fmt):
    """register() writes into docling's lru_cached factory under the exact call the
    pipeline makes. If a docling upgrade changes that call, this fails here, by
    name, instead of every Word file failing in production."""
    from docling.datamodel.base_models import InputFormat
    from docling.document_converter import DocumentConverter

    conv = DocumentConverter(format_options=pd.office_format_options())
    pipeline = conv._get_pipeline(InputFormat[fmt])  # raises if 'ragp_vlm' is unknown
    assert any(isinstance(m, pd.VlmPictureDescriptionModel) for m in pipeline.enrichment_pipe)


def test_switched_off_leaves_docling_defaults(monkeypatch):
    monkeypatch.setenv("DOCLING_DESCRIBE_PICTURES", "false")
    assert pd.office_format_options() == {}


def test_many_pictures_are_not_cut_by_the_stall_watch(tmp_path, monkeypatch):
    """20 pictures on one page, 1 s each, 4 at a time: ~5 s with no docling page
    handoff. The warm worker's stall watch (4 s here) must see that as progress."""
    import docx
    from docx.shared import Inches

    from app.workers import worker as worker_mod
    from tests import _warm_targets as targets

    d = docx.Document()
    for _ in range(20):
        d.add_picture(_png(200), width=Inches(0.5))
    path = tmp_path / "many.docx"
    d.save(str(path))

    monkeypatch.setenv("DOCLING_STALL_SECONDS", "4")
    monkeypatch.setenv("DOCLING_PICTURE_CONCURRENCY", "4")
    monkeypatch.setattr(worker_mod, "_warm_docling_loop", targets.slow_vlm_docling_loop)
    w = worker_mod._WarmDoclingWorker()
    try:
        result = w.convert(str(path), "c1", False, False, 120)
    finally:
        w._kill()
    assert result["markdown"].count("תיאור") == 20


def _annotate(describe, n: int, concurrency: int) -> tuple[list[str], float]:
    """Run the engine's picture step over n stand-in pictures with `describe`."""
    import time
    from types import SimpleNamespace

    fake = SimpleNamespace(options=SimpleNamespace(concurrency=concurrency), _describe=describe)
    t0 = time.monotonic()
    out = list(pd.VlmPictureDescriptionModel._annotate_images(fake, list(range(n))))
    return out, time.monotonic() - t0


def test_a_hung_picture_is_left_undescribed_and_the_document_goes_on(monkeypatch):
    """A VLM request that never returns must not hold the conversion: the beat keeps
    the stall watch quiet, so the picture step needs its own deadline."""
    import threading

    monkeypatch.setenv("DOCLING_PICTURE_TIMEOUT_SECONDS", "2")
    hang = threading.Event()

    def describe(i):
        if i == 1:
            hang.wait(30)
        return f"pic {i}"

    try:
        out, secs = _annotate(describe, 4, concurrency=2)
    finally:
        hang.set()
    assert out == ["pic 0", "", "pic 2", "pic 3"]
    assert secs < 10


def test_pictures_queued_behind_hung_threads_are_not_waited_for(monkeypatch):
    """Every thread held by a hung picture: the rest can never start, so they are
    left undescribed rather than waited on forever."""
    import threading

    monkeypatch.setenv("DOCLING_PICTURE_TIMEOUT_SECONDS", "1")
    hang = threading.Event()

    def describe(i):
        hang.wait(30)
        return f"pic {i}"

    try:
        out, secs = _annotate(describe, 5, concurrency=2)
    finally:
        hang.set()
    assert out == [""] * 5
    assert secs < 10


def test_a_hung_vlm_fallback_is_cut_by_the_stall_watch(tmp_path, monkeypatch):
    """A PDF whose text layer sends it to the VLM, and a VLM request that never
    returns: the warm worker must be cut as stalled, not held for the 2 h ceiling
    (2026-10-31 local soak: h2504.pdf held a worker for 31 min)."""
    import time

    import fitz

    from app.workers import worker as worker_mod
    from tests import _warm_targets as targets

    path = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "hello")
    doc.save(str(path))

    monkeypatch.setenv("DOCLING_STALL_SECONDS", "5")
    monkeypatch.setattr(worker_mod, "_warm_docling_loop", targets.hung_vlm_fallback_loop)
    w = worker_mod._WarmDoclingWorker()
    t0 = time.monotonic()
    try:
        with pytest.raises(TimeoutError, match="stalled"):
            w.convert(str(path), "c1", False, False, 120)
    finally:
        w._kill()
    assert time.monotonic() - t0 < 60
