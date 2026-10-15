"""Slide-deck detection and the slides converter (VLM mocked — no network)."""
from __future__ import annotations

import multiprocessing
from pathlib import Path
from unittest.mock import MagicMock, patch

import pymupdf as fitz
import pytest

from app.converters.slides import SlidesConverter, is_slide_deck
from app.workers import worker


def _pdf(path: Path, size, pages, *, text="", words_per_page=None, image_frac=0.0,
         creator="", producer="") -> Path:
    """Write a PDF of `pages` pages at `size` (w, h). Each page gets `text` (or
    `words_per_page` words), and an image covering `image_frac` of the page."""
    w, h = size
    doc = fitz.open()
    for _ in range(pages):
        page = doc.new_page(width=w, height=h)
        body = text or " ".join(f"w{i}" for i in range(words_per_page or 0))
        if body:
            page.insert_text((40, 60), body, fontsize=10)
        if image_frac:
            pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20), False)
            pix.clear_with(128)
            side = image_frac ** 0.5
            page.insert_image(fitz.Rect(0, 0, w * side, h * side), pixmap=pix)
    doc.set_metadata({"creator": creator, "producer": producer})
    doc.save(str(path))
    doc.close()
    return path


# ------------------------------------------------------------------ detection

def test_powerpoint_metadata_landscape(tmp_path):
    p = _pdf(tmp_path / "a.pdf", (960, 540), 3, words_per_page=100, creator="Microsoft® PowerPoint® LTSC")
    assert is_slide_deck(p) == "metadata"


def test_google_slides_metadata(tmp_path):
    p = _pdf(tmp_path / "a.pdf", (960, 540), 2, words_per_page=100, creator="Google")
    assert is_slide_deck(p) == "metadata"


def test_metadata_needs_landscape(tmp_path):
    # A portrait Google Docs export is a document, whatever the producer says
    p = _pdf(tmp_path / "a.pdf", (595, 842), 2, words_per_page=100, creator="Google")
    assert is_slide_deck(p) is None


def test_uniform_16_9_pages_without_metadata(tmp_path):
    p = _pdf(tmp_path / "a.pdf", (1920, 1080), 4, words_per_page=100)
    assert is_slide_deck(p) == "aspect"


def test_4_3_deck_needs_metadata(tmp_path):
    p = _pdf(tmp_path / "a.pdf", (720, 540), 4, words_per_page=100)
    assert is_slide_deck(p) is None


def test_landscape_a4_report_stays_on_docling(tmp_path):
    p = _pdf(tmp_path / "a.pdf", (842, 595), 5, words_per_page=200, producer="FPDF 1.7")
    assert is_slide_deck(p) is None


def test_portrait_report_stays_on_docling(tmp_path):
    p = _pdf(tmp_path / "a.pdf", (595, 842), 5, words_per_page=200)
    assert is_slide_deck(p) is None


def test_visual_pages_with_stripped_metadata(tmp_path):
    # landscape Letter, big picture + a title on most pages: a deck printed to PDF
    p = _pdf(tmp_path / "a.pdf", (792, 612), 4, words_per_page=5, image_frac=0.6)
    assert is_slide_deck(p) == "visual"


def test_visual_rule_needs_majority(tmp_path):
    doc = fitz.open()
    for i in range(4):
        page = doc.new_page(width=792, height=612)
        if i == 0:
            pix = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 20, 20), False)
            page.insert_image(fitz.Rect(0, 0, 700, 500), pixmap=pix)
            page.insert_text((40, 60), "title", fontsize=10)
        else:
            page.insert_text((40, 60), " ".join(f"w{k}" for k in range(200)), fontsize=10)
    p = tmp_path / "a.pdf"
    doc.save(str(p))
    assert is_slide_deck(p) is None


def test_scanned_portrait_is_not_a_deck(tmp_path):
    p = _pdf(tmp_path / "a.pdf", (595, 842), 3, image_frac=1.0)
    assert is_slide_deck(p) is None


# ------------------------------------------------------------------ converter

def _convert(pdf: Path, outputs: list[str]):
    conv = SlidesConverter(vlm_concurrency=2)
    vlm = MagicMock()
    vlm.transcribe_page.side_effect = list(outputs)
    with patch.object(SlidesConverter, "_render", return_value="IMG"), \
         patch.object(conv, "_ensure_vlm", side_effect=lambda: setattr(conv, "_vlm", vlm)):
        return conv.convert(pdf), vlm


def test_every_slide_goes_to_the_vlm_in_order(tmp_path):
    p = _pdf(tmp_path / "a.pdf", (960, 540), 3, words_per_page=20, creator="Keynote")
    out, vlm = _convert(p, ["## one", "## two", "## three"])
    assert vlm.transcribe_page.call_count == 3
    assert out.split("\n\n---\n\n") == [
        "<!-- page 1 -->\n\n## one", "<!-- page 2 -->\n\n## two", "<!-- page 3 -->\n\n## three",
    ]


def test_number_not_on_text_slide_is_tagged_not_retried(tmp_path):
    p = _pdf(tmp_path / "a.pdf", (960, 540), 1, text="revenue 1,500", creator="Keynote")
    out, vlm = _convert(p, ["revenue 1,500 and 99,000"])
    assert vlm.transcribe_page.call_count == 1
    assert "<!-- slides: numbers not on the slide: 99,000 -->" in out


def test_number_check_skipped_on_slide_with_figure(tmp_path):
    # map labels live in the raster: the text layer cannot vouch for them
    p = _pdf(tmp_path / "a.pdf", (960, 540), 1, text="map", image_frac=0.5, creator="Keynote")
    out, _ = _convert(p, ["block 6675 and 8921"])
    assert "numbers not on the slide" not in out


# ------------------------------------------------------------------ worker reroute

def _record_target(file_path, converter_type, do_ocr, result_queue, do_table_structure=True):
    result_queue.put(("success", {"markdown": "# x", "num_pages": 1, "doc_name": "x",
                                  "doc_origin": f"converter:{converter_type}",
                                  "converter_used": converter_type}))


@pytest.mark.parametrize("size,creator,expected", [
    ((960, 540), "Microsoft PowerPoint", "converter:slides"),
    ((595, 842), "", "converter:dbank"),
])
def test_dbank_lane_reroutes_detected_decks(tmp_path, size, creator, expected):
    p = _pdf(tmp_path / "a.pdf", size, 2, words_per_page=50, creator=creator)
    result = worker._run_converter_with_timeout(
        str(p), "cid", "dbank", timeout_seconds=30, _target=_record_target)
    assert result["doc_origin"] == expected


@pytest.mark.parametrize("lane", ["docling", "vlm"])
def test_other_converters_are_not_rerouted(tmp_path, lane):
    p = _pdf(tmp_path / "a.pdf", (960, 540), 2, words_per_page=50, creator="Keynote")
    result = worker._run_converter_with_timeout(
        str(p), "cid", lane, timeout_seconds=30, _target=_record_target)
    assert result["doc_origin"] == f"converter:{lane}"


def test_precheck_fault_falls_through_to_dbank(tmp_path):
    result = worker._run_converter_with_timeout(
        str(tmp_path / "missing.pdf"), "cid", "dbank", timeout_seconds=30, _target=_record_target)
    assert result["doc_origin"] == "converter:dbank"


# ------------------------------------------------------------------ dbank: non-PDF files

def _write_docx(path: Path) -> Path:
    import docx
    d = docx.Document()
    d.add_paragraph("בנקאות פתוחה בישראל מאפשרת ללקוחות גישה למידע פיננסי " * 12)
    d.save(str(path))
    return path


@pytest.mark.parametrize("name,maker", [
    ("a.docx", _write_docx),
    ("a.html", lambda p: (p.write_text("<html><body><p>שלום עולם</p></body></html>"), p)[1]),
    ("a.csv", lambda p: (p.write_text("name,amount\nalpha,100\n"), p)[1]),
    ("a.txt", lambda p: (p.write_text("hello " * 200), p)[1]),
])
def test_dbank_sends_non_pdf_to_docling(tmp_path, name, maker):
    p = maker(tmp_path / name)
    result = worker._run_converter_with_timeout(
        str(p), "cid", "dbank", timeout_seconds=30, _target=_record_target)
    assert result["doc_origin"] == "converter:docling"


def test_dbank_keeps_pdf_and_images(tmp_path):
    from PIL import Image
    pdf = _pdf(tmp_path / "r.pdf", (595, 842), 1, words_per_page=50)
    png = tmp_path / "i.png"
    Image.new("RGB", (40, 40), "white").save(png)
    for p in (pdf, png):
        result = worker._run_converter_with_timeout(
            str(p), "cid", "dbank", timeout_seconds=30, _target=_record_target)
        assert result["doc_origin"] == "converter:dbank"


def test_dbank_pdf_named_without_extension_stays_on_dbank(tmp_path):
    # detection is by content: a stored file key may carry no .pdf suffix
    p = _pdf(tmp_path / "r.pdf", (595, 842), 1, words_per_page=50)
    bare = tmp_path / "blob"
    bare.write_bytes(p.read_bytes())
    result = worker._run_converter_with_timeout(
        str(bare), "cid", "dbank", timeout_seconds=30, _target=_record_target)
    assert result["doc_origin"] == "converter:dbank"
